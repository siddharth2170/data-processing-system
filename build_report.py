from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "Concurrent_Data_Processing_System_APA7_Report.docx"


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(text, first_indent=True):
    p = doc.add_paragraph(style="Body Text")
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    p.add_run(text)
    return p


def add_heading(text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_figure(image_name, caption, width=6.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(ROOT / "assets" / image_name), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = False
    label, title = caption.split(". ", 1)
    r = cap.add_run(label + ". ")
    r.bold = True
    r2 = cap.add_run(title)
    r2.italic = True


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.page_width = Inches(8.5)
section.page_height = Inches(11)
add_page_number(section.header.paragraphs[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)

body = styles["Body Text"]
body.font.name = "Times New Roman"
body._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
body._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
body.font.size = Pt(12)
body.paragraph_format.line_spacing = 2
body.paragraph_format.space_after = Pt(0)
body.paragraph_format.widow_control = True

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True

styles["Title"].font.size = Pt(15)
styles["Heading 1"].font.size = Pt(12)
styles["Heading 2"].font.size = Pt(12)
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(0)
styles["Heading 1"].paragraph_format.keep_with_next = True
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(0)
styles["Heading 2"].paragraph_format.keep_with_next = True

if "Reference" not in styles:
    ref_style = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    ref_style.base_style = styles["Normal"]
else:
    ref_style = styles["Reference"]
ref_style.paragraph_format.left_indent = Inches(0.5)
ref_style.paragraph_format.first_line_indent = Inches(-0.5)
ref_style.paragraph_format.line_spacing = 2
ref_style.paragraph_format.space_after = Pt(0)

caption_style = styles["Caption"]
caption_style.font.name = "Times New Roman"
caption_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
caption_style.font.size = Pt(10)
caption_style.font.color.rgb = RGBColor(0, 0, 0)
caption_style.paragraph_format.space_after = Pt(0)

# APA student-style title page
for _ in range(5):
    doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Concurrent Data Processing System in Java and Go")
for line in ("Siddharth Malhotra", "Data Processing System Assignment", "October 8, 2026"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2
    p.add_run(line)

doc.add_page_break()

title2 = doc.add_paragraph(style="Title")
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
title2.add_run("Concurrent Data Processing System in Java and Go")

add_body(
    "This project compares equivalent three-worker systems in Java and Go. Each processes 10 queued integer "
    "tasks, records nine successes and one intentional input failure, and sorts the saved results for verification."
)

add_heading("Java Concurrency and Exception Handling", 1)
add_body(
    "The Java TaskQueue wraps an ArrayDeque with addTask and getTask methods. A ReentrantLock protects queue "
    "state, while a Condition suspends workers when no task is available. Unlock runs in finally blocks, so an "
    "exception cannot leave the queue locked. When the producer closes the queue, signalAll wakes waiting "
    "workers; getTask returns null only after the closed queue is drained. This creates a termination condition "
    "without polling or poison-pill tasks."
)
add_body(
    "A fixed ExecutorService owns the workers. Main calls shutdown after submission and awaits termination with "
    "a timeout. Oracle (2026a) states that orderly shutdown rejects new work while allowing submitted work to "
    "finish. A synchronized list protects results. Worker try-catch blocks convert invalid inputs into ERROR "
    "records, preserve InterruptedException status, and continue processing when possible. Try-with-resources "
    "closes the output writer, and IOException is logged."
)

add_heading("Go Concurrency and Error Handling", 1)
add_body(
    "Go uses a buffered channel as its TaskQueue. AddTask sends a value, GetTask exposes a receive-only channel, "
    "and Close closes the channel after submission. A goroutine ranges over the channel and stops after it is "
    "closed and drained. Closing a channel indicates that no more values will be sent (The Go Authors, 2026a). "
    "A sync.WaitGroup prevents output from being written before every worker returns."
)
add_body(
    "Go functions return errors rather than throwing exceptions. processTask returns a Result and an error; a "
    "bad value becomes a logged ERROR record without stopping other goroutines. A mutex protects the result "
    "slice, and deferred Unlock calls limit the critical section to slice operations. The output function wraps "
    "create, write, flush, and close errors. Deferred cleanup runs on every return path."
)

add_heading("Comparison of the Concurrency Models", 1)
add_body(
    "Java coordinates shared state with locks and a condition; ExecutorService separately manages thread "
    "lifecycle. ReentrantLock provides monitor-like exclusion, but lock release remains the programmer's "
    "responsibility, which is why Oracle (2026b) recommends try-finally. Go moves task ownership through a "
    "channel. Closing the channel communicates completion, while WaitGroup provides the join operation. A mutex "
    "is still required for the shared result slice."
)
add_body(
    "Go needs less queue synchronization code because send, receive, and close express the handoff directly. "
    "Its pipeline guidance recommends closing outbound channels and draining receive loops (The Go Authors, "
    "2014). Java gives more explicit control over waiting and interruption but requires stricter lock discipline."
)

add_heading("Deadlock Prevention and Verification", 1)
add_body(
    "Neither implementation nests locks or performs work while holding one. Java unlocks in finally blocks and "
    "wakes all waiters at closure; Go closes the task channel once after all sends. Both saved 10 records with "
    "nine OK and one ERROR status. The Go race-enabled test produced no race report."
)

doc.add_page_break()
add_heading("References", 1)
refs = [
    "Oracle. (2026a). ExecutorService (Java SE 25 & JDK 25). https://docs.oracle.com/en/java/javase/25/docs/api/java.base/java/util/concurrent/ExecutorService.html",
    "Oracle. (2026b). ReentrantLock (Java SE 25 & JDK 25). https://docs.oracle.com/en/java/javase/25/docs/api/java.base/java/util/concurrent/locks/ReentrantLock.html",
    "The Go Authors. (2014, March 13). Go concurrency patterns: Pipelines and cancellation. https://go.dev/blog/pipelines",
    "The Go Authors. (2026a). The Go programming language specification: Channel types. https://go.dev/ref/spec#Channel_types",
]
for ref in refs:
    doc.add_paragraph(ref, style="Reference")

add_heading("Repository", 1)
p = doc.add_paragraph(style="Body Text")
p.paragraph_format.first_line_indent = Inches(0.5)
p.add_run("The publication-ready repository is located in the submitted data-processing-system folder. GitHub publishing requires re-authentication for the configured account; insert the resulting repository URL before final submission.")

doc.add_page_break()
add_heading("Appendix A", 1)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Java Source Code and Sample Output")
r.bold = True
add_figure("java-code.png", "Figure A1. Java synchronized queue and worker-pool code excerpt")

doc.add_page_break()
add_heading("Appendix A Continued", 1)
add_figure("java-output.png", "Figure A2. Java execution log showing worker lifecycle and final counts")

doc.add_page_break()
add_heading("Appendix B", 1)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Go Source Code and Sample Output")
r.bold = True
add_figure("go-code.png", "Figure B1. Go channel queue and synchronized result-store code excerpt")

doc.add_page_break()
add_heading("Appendix B Continued", 1)
add_figure("go-output.png", "Figure B2. Go execution log showing worker lifecycle and final counts")

# Keep all body runs in the required font.
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//w:drawing"):
            continue
        size = run.font.size.pt if run.font.size else 12
        set_font(run, size=size, bold=run.bold, italic=run.italic)

doc.core_properties.title = "Concurrent Data Processing System in Java and Go"
doc.core_properties.author = "Siddharth Malhotra"
doc.core_properties.subject = "Concurrency and exception handling comparison"
doc.save(OUTPUT)
print(OUTPUT)
